# -*- coding: utf-8 -*-
"""
Reproduce todas las cifras de Memoria/analisis/capacidad-acceso-conexion.md.

Cruza el mapa de capacidad de acceso de generacion de i-DE (R1-001) con el del
gestor de la red de transporte (REE), usando la columna «Nudo Afeccion RdT» del
primero como clave sobre el segundo, y georreferencia los puntos de conexion
contra la parcela candidata del TFM.kmz.

Uso:  python capacidad_acceso.py
"""
import io
import math
import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
KMZ = os.path.join(RAIZ, "Datos", "Google Earth", "TFM.kmz")
CAP = os.path.join(RAIZ, "Datos", "Capacidad AyC Generacion")
F_IDE = os.path.join(CAP, "2026_07_01_R1-001_Generacion.xlsx")
F_REE = os.path.join(CAP, "2026_08_03_GRT_generacion.xlsx")

NS = {"k": "http://www.opengis.net/kml/2.2"}
KML = "{http://www.opengis.net/kml/2.2}"

# Columnas del fichero REE (1-indexadas), segun la hoja «Leyenda»
COL_REE = {
    "nombre": 1, "cod": 2, "ccaa": 3,
    "nodisp_mges": 48, "nodisp_mpe": 49, "motivo_nodisp": 52,
    "disp_mges_rdt": 53, "disp_mpe_rdt": 54,
    "disp_mges_rdd": 55, "disp_mpe_rdd": 56,
    "lib_mges": 61, "lib_mpe": 62,
    "otor_mges": 63, "otor_mpe": 64,
    "dispauto_mges": 65, "dispauto_mpe": 66,
}


def num(v):
    """Convierte a float admitiendo el formato numerico espanol de los ficheros."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def a_utm30n(lon, lat):
    """WGS84 -> UTM 30N. Formulacion estandar, precision centimetrica."""
    a, f = 6378137.0, 1 / 298.257223563
    e2 = f * (2 - f)
    ep2 = e2 / (1 - e2)
    k0 = 0.9996
    lon0 = math.radians(-3.0)  # meridiano central del huso 30
    lam, phi = math.radians(lon), math.radians(lat)
    N = a / math.sqrt(1 - e2 * math.sin(phi) ** 2)
    T = math.tan(phi) ** 2
    C = ep2 * math.cos(phi) ** 2
    A = (lam - lon0) * math.cos(phi)
    e4, e6 = e2 * e2, e2 ** 3
    M = a * ((1 - e2 / 4 - 3 * e4 / 64 - 5 * e6 / 256) * phi
             - (3 * e2 / 8 + 3 * e4 / 32 + 45 * e6 / 1024) * math.sin(2 * phi)
             + (15 * e4 / 256 + 45 * e6 / 1024) * math.sin(4 * phi)
             - (35 * e6 / 3072) * math.sin(6 * phi))
    x = k0 * N * (A + (1 - T + C) * A ** 3 / 6
                  + (5 - 18 * T + T * T + 72 * C - 58 * ep2) * A ** 5 / 120) + 500000
    y = k0 * (M + N * math.tan(phi) * (A * A / 2
              + (5 - T + 9 * C + 4 * C * C) * A ** 4 / 24
              + (61 - 58 * T + T * T + 600 * C - 330 * ep2) * A ** 6 / 720))
    return x, y


def geometria():
    """Centroide de la parcela (Layout1) y punto del CPD, en UTM 30N."""
    with zipfile.ZipFile(KMZ) as z:
        kml = [n for n in z.namelist() if n.endswith(".kml")][0]
        raiz = ET.fromstring(z.read(kml))
    puntos = {}
    for pm in raiz.iter(KML + "Placemark"):
        n = pm.find("k:name", NS)
        nombre = n.text.strip() if n is not None and n.text else ""
        for tipo in ("Point", "Polygon"):
            for geo in pm.iter(KML + tipo):
                c = geo.find(".//k:coordinates", NS)
                if c is None or not c.text:
                    continue
                pts = [(float(t.split(",")[0]), float(t.split(",")[1]))
                       for t in c.text.split() if "," in t]
                puntos.setdefault((nombre, tipo), pts)
    lay = puntos[("Layout1", "Polygon")]
    cpd = puntos[("CPD", "Point")][0]
    parc = a_utm30n(sum(p[0] for p in lay) / len(lay),
                    sum(p[1] for p in lay) / len(lay))
    return parc, a_utm30n(*cpd)


def carga_ree():
    wb = openpyxl.load_workbook(F_REE, data_only=True)
    ws = wb["Ficha Capacidades"]
    nudos = {}
    for r in range(4, ws.max_row + 1):
        nombre = ws.cell(r, COL_REE["nombre"]).value
        if not nombre:
            continue
        nudos[str(nombre).strip().upper()] = {
            k: ws.cell(r, c).value for k, c in COL_REE.items()
        }
    wb.close()
    return nudos


def carga_ide():
    wb = openpyxl.load_workbook(F_IDE, read_only=True, data_only=True)
    filas = list(wb["Datos"].iter_rows(values_only=True))
    wb.close()
    return filas[1:]


def main():
    (px, py), (cx, cy) = geometria()
    print(f"Parcela  UTM30N X={px:.0f} Y={py:.0f}")
    print(f"CPD      UTM30N X={cx:.0f} Y={cy:.0f}")
    print(f"Distancia entre centroides: {math.dist((cx, cy), (px, py)) / 1000:.2f} km\n")

    ree = carga_ree()
    ide = carga_ide()
    print(f"Nudos REE: {len(ree)}   |   Puntos i-DE: {len(ide)}\n")

    # --- 1. Verificacion de la regla del 10 % de la DA 10.a del RDL 7/2026 ---
    con_motivo, sin_motivo, ratios = 0, 0, []
    for v in ree.values():
        lib = num(v["lib_mpe"]) or 0
        if lib <= 0:
            continue
        if "Concurso por resoluci" in str(v["motivo_nodisp"] or ""):
            con_motivo += 1
        else:
            sin_motivo += 1
        nd = num(v["nodisp_mpe"]) or 0
        if nd > 0:
            ratios.append(lib / nd)
    ratios.sort()
    print("### Verificacion de la DA 10.a (liberacion del 10 % en nudos de concurso)")
    print(f"  Nudos con capacidad liberada MPE > 0            : {con_motivo + sin_motivo}")
    print(f"    de ellos con motivo 'Concurso por resolucion' : {con_motivo}")
    print(f"    sin ese motivo (deberia ser 0)                : {sin_motivo}")
    print(f"  Ratio liberada/no-disponible: mediana={ratios[len(ratios) // 2]:.3f} "
          f"min={min(ratios):.3f} max={max(ratios):.3f}\n")

    # --- 2. Los dos nudos de 220 kV del entorno ---
    print("### Nudos de 220 kV del entorno del CPD")
    for nombre in ("TRES CANTOS 220", "TRES CANTOS GIS 220"):
        v = ree[nombre]
        print(f"  {nombre} (cod {v['cod']})")
        print(f"     margen/no disponible MPE : {v['nodisp_mpe']} MW")
        print(f"     motivo                   : {str(v['motivo_nodisp'] or '-').strip()}")
        print(f"     disponible MPE via RdD   : {v['disp_mpe_rdd']} MW")
        print(f"     AUTOCONSUMO RDL 7/2026   : liberada={v['lib_mpe']} "
              f"otorgada={v['otor_mpe']} disponible={v['dispauto_mpe']} MW")
    print()

    # --- 3. Puntos de i-DE proximos y su nudo de afeccion ---
    print("### Puntos de conexion i-DE a <= 6 km de la parcela")
    cand = []
    for r in ide:
        x, y = num(r[3]), num(r[4])
        if x is None or y is None:
            continue
        d = math.dist((x, y), (px, py)) / 1000
        if d <= 6:
            cand.append((d, r))
    cand.sort(key=lambda t: t[0])
    for d, r in cand:
        afec = str(r[13]).strip().upper() if r[13] else "(sin afeccion)"
        v = ree.get(afec)
        disp_ree = v["disp_mpe_rdd"] if v else "?"
        lib_ree = v["dispauto_mpe"] if v else "?"
        print(f"  {d:5.2f} km  {str(r[17]).strip():24} {str(r[6]):>3} kV  "
              f"i-DE disp={str(r[7]):>6} MW  ->  {afec}: "
              f"REE disp={disp_ree} MW, autoconsumo={lib_ree} MW")

    # --- 4. Recuento de nudos de afeccion en Madrid ---
    c = Counter(str(r[13]).strip() for r in ide if r[1] and "Madrid" in str(r[1]))
    print(f"\n### Nudo de afeccion 'TRES CANTOS 220' en Madrid: "
          f"{c.get('TRES CANTOS 220', 0)} puntos de conexion")

    # --- 5. Techo de potencia por la condicion b) ---
    print("\n### Condicion b) de la DA 10.a: P1_contratada / P_generacion >= 0,5")
    for etiqueta, p1 in (("CUPS actual del complejo (factura i-DE 08/2025)", 2300),
                         ("punto propio del CPD (estimado)", 600)):
        print(f"  P1 = {p1:>5} kW  ->  P_gen maxima = {p1 / 0.5:.0f} kW "
              f"= {p1 / 0.5 / 1000:.1f} MW   [{etiqueta}]")


if __name__ == "__main__":
    main()
